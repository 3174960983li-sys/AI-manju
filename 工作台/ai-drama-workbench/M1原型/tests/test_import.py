from io import BytesIO
import sqlite3
from zipfile import ZipFile

import pytest
from docx import Document

from manju_m1.db import Database
from manju_m1.scripts import MAX_UPLOAD_BYTES, ScriptReadError, parse_script
from manju_m1.service import ProjectService


@pytest.fixture
def service(tmp_path):
    db = Database(tmp_path / "test.sqlite3")
    db.initialize()
    return ProjectService(db)


@pytest.mark.parametrize("encoding", ["utf-8-sig", "gb18030", "utf-16"])
def test_txt_encodings(encoding):
    parsed = parse_script("剧本.txt", "第一场\r\n林禾：你好。".encode(encoding))
    assert parsed.text == "第一场\n林禾：你好。"
    assert parsed.encoding == encoding


@pytest.mark.parametrize("name, data", [("empty.txt", b""), ("blank.txt", b" \n\t"), ("binary.txt", b"a\x00b"), ("old.doc", b"abc"), ("broken.docx", b"not a zip"), ("large.txt", b"a" * (MAX_UPLOAD_BYTES + 1))])
def test_reject_invalid_files(name, data):
    with pytest.raises(ScriptReadError):
        parse_script(name, data)


def docx_bytes():
    doc = Document()
    doc.add_paragraph("第一场 内景 门厅 夜")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "林禾"
    table.cell(0, 1).text = "谁在外面？"
    doc.add_paragraph("门外无人回答。")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_preserves_paragraph_table_order():
    parsed = parse_script("剧本.DOCX", docx_bytes())
    assert parsed.text == "第一场 内景 门厅 夜\n林禾\t谁在外面？\n门外无人回答。"
    assert "文本框" in parsed.warning


def test_reject_zip_without_document():
    buf = BytesIO()
    with ZipFile(buf, "w") as archive:
        archive.writestr("fake.txt", "abc")
    with pytest.raises(ScriptReadError, match="DOCX正文"):
        parse_script("wrong.docx", buf.getvalue())


def test_docx_expansion_limit(monkeypatch):
    monkeypatch.setattr("manju_m1.scripts.MAX_DOCX_EXPANDED_BYTES", 1)
    with pytest.raises(ScriptReadError, match="解压后"):
        parse_script("script.docx", docx_bytes())


def test_create_import_restart_and_original_bytes(service):
    pid = service.create_project("  雨夜来信  ", "项目简介")
    data = "第一场\n主角开门。".encode("utf-8")
    sid, created = service.import_script(pid, "..\\..\\剧本.txt", data)
    assert created
    db = Database(service.db.path)
    db.initialize()
    assert db.get_project(pid)["name"] == "雨夜来信"
    stored = db.get_script(pid, sid)
    assert stored["content"] == data.decode("utf-8")
    assert stored["original_bytes"] == data
    assert stored["filename"] == "剧本.txt"
    assert db.list_scripts(pid)[0]["char_count"] == len(stored["content"])


def test_duplicate_is_per_project_and_does_not_overwrite(service):
    p1 = service.create_project("A")
    p2 = service.create_project("B")
    data = b"scene 1"
    first, _ = service.import_script(p1, "one.txt", data)
    again, created = service.import_script(p1, "renamed.txt", data)
    other, other_created = service.import_script(p2, "one.txt", data)
    assert again == first and not created
    assert other != first and other_created
    assert len(service.db.list_scripts(p1)) == 1
    assert service.db.get_script(p1, other) is None


def test_invalid_input_leaves_database_unchanged(service):
    with pytest.raises(ValueError):
        service.create_project("  ")
    assert service.db.list_projects() == []
    pid = service.create_project("真实项目")
    with pytest.raises(ScriptReadError):
        service.import_script(pid, "broken.docx", b"broken")
    assert service.db.list_scripts(pid) == []
    with pytest.raises(ValueError, match="项目不存在"):
        service.import_script("unknown", "a.txt", b"hello")
    with service.db.connection() as connection:
        with pytest.raises(sqlite3.IntegrityError):
            connection.execute("INSERT INTO scripts VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", ("bad", "missing", "a.txt", "txt", "utf-8", "a", b"a", "abc", "", "now"))


def test_new_upload_is_not_an_overwrite(service):
    pid = service.create_project("版本测试")
    first, _ = service.import_script(pid, "剧本.txt", b"first")
    second, _ = service.import_script(pid, "剧本.txt", b"second")
    assert first != second
    assert len(service.db.list_scripts(pid)) == 2
    assert service.db.get_script(pid, first)["content"] == "first"