from pathlib import Path
import json,hashlib,shutil
from docx import Document
base=Path(r'E:\AI漫剧');work=Path(__file__).parent
names=['project-script-analysis','photoreal-drama-director','shot-language-design','video-prompt-authoring','tig-blocking-map']
source=base/'02_制作方法/02_专业技能/自研Skill'
backup=base/'05_文件整理/03_停用归档/Skill历史版本/2026-10-27_五Skill全面翻新前'
assert not backup.exists(), 'Backup must be a new snapshot'
backup.mkdir(parents=True)
manifest={};inventory=[]
for name in names:
    shutil.copytree(source/name,backup/'自研Skill'/name)
    for f in sorted((source/name).rglob('*')):
        if f.is_file():
            rel=f.relative_to(source);h=hashlib.sha256(f.read_bytes()).hexdigest()
            assert hashlib.sha256((backup/'自研Skill'/rel).read_bytes()).hexdigest()==h
            manifest['自研Skill/'+str(rel)]=h
            inventory.append({'file':str(f),'size':f.stat().st_size,'sha256':h,'active':('试跑' not in f.parts and '原始附件' not in f.parts)})
for dirname in ['01_流程规范','02_专业技能/Agent职责与交接']:
    src=base/'02_制作方法'/dirname;dst=backup/'制作方法'/dirname
    shutil.copytree(src,dst)
for name in names:
    installed=Path(r'C:\Users\Administrator\.codex\skills')/name
    if installed.exists():
        shutil.copytree(installed,backup/'安装副本'/name)
        files=[{'path':str(f.relative_to(installed)),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in installed.rglob('*') if f.is_file()]
        (work/(name+'-installed-before.json')).write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
(backup/'备份校验.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(work/'before-inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
notes=base/'03_参考资料/02_知识笔记/书籍知识文档'
for sub in ['角色表演','运镜设计']:
    texts=[]
    for f in sorted((notes/sub).glob('*.docx')):
        d=Document(f);lines=[p.text for p in d.paragraphs if p.text.strip()]
        lines += [' | '.join(c.text for c in row.cells) for table in d.tables for row in table.rows]
        texts.append('\n# '+f.name+'\n'+'\n'.join(lines))
    (work/(sub+'教程正文.txt')).write_text('\n'.join(texts),encoding='utf-8')
print('Full 5-package backup verified:',len(manifest),'files')
print('Active files:',sum(x['active'] for x in inventory))
print('Extracted tutorial text:',[(p.name,len(p.read_text(encoding='utf-8'))) for p in work.glob('*教程正文.txt')])
