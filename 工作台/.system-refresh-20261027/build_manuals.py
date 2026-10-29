from pathlib import Path
import re,shutil,json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
base=Path(r'E:\AI漫剧');work=Path(__file__).parent
root=base/'02_制作方法/02_专业技能/自研Skill'
manual=base/'02_制作方法/04_通用模板/分镜运镜与表演操作手册'
out=manual/'Word阅读版';out.mkdir(exist_ok=True)
items=[('分镜设计操作指南','shot-language-design'),('运镜设计操作指南','shot-language-design'),('角色表演与微表情设计操作手册','photoreal-drama-director')]
def inline(par,text):
    for token in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)',text):
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if m:
            label,target=m.groups()
            if '://' not in target:par.add_run(label);continue
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),par.part.relate_to(target,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');p=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');p.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');p.append(u);r.append(p);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
        elif token.startswith('**') and token.endswith('**'):par.add_run(token[2:-2]).bold=True
        elif token.startswith('`') and token.endswith('`'):par.add_run(token[1:-1])
        else:par.add_run(token)
def rowcells(line):return [x.strip() for x in line.strip().strip('|').split('|')]
manifest=[]
for title,owner in items:
    src=root/owner/'references'/(title+'.md')
    text=src.read_text(encoding='utf-8-sig').replace('声音/body协同','声音和身体协同').replace('沿门侧短移，末段收至双人中景','沿门侧斜向前移，末段收至双人中景')
    src.write_text(text,encoding='utf-8');(work/(title+'.md')).write_text(text,encoding='utf-8')
    shutil.copyfile(src,manual/'MD同版导出'/(title+'.md'))
    doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.9);s.left_margin=s.right_margin=Cm(2)
    for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        st=doc.styles[n];st.font.name='Microsoft YaHei';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor(0,0,0)
        for b in list(st.element.xpath('.//w:pBdr')):b.getparent().remove(b)
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.15;doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    for n,size in [('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        doc.styles[n].font.size=Pt(size);doc.styles[n].paragraph_format.space_before=Pt(10);doc.styles[n].paragraph_format.space_after=Pt(6)
    s.header.paragraphs[0].text=title+'  V1.0';s.header.paragraphs[0].runs[0].font.size=Pt(8)
    footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('第 ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f);footer.add_run(' 页')
    lines=text.splitlines();i=0;covered=[]
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|\s*:?-',lines[i+1]):
            header=rowcells(line);rows=[];i+=2
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=rowcells(lines[i]);assert len(cells)==len(header),(title,cells);rows.append(cells);i+=1
            if len(header)>4:
                for cells in rows:
                    doc.add_heading(cells[0],3);covered.append(cells[0])
                    for field_index,(label,value) in enumerate(zip(header[1:],cells[1:])):
                        p=doc.add_paragraph();p.add_run(label+'：').bold=True;inline(p,value);covered.append(value)
                        p.paragraph_format.keep_with_next=field_index<len(cells)-2
            else:
                table=doc.add_table(rows=1,cols=len(header));table.style='Table Grid';table.autofit=False
                widths=[17/len(header)]*len(header)
                if len(header)==4:widths=[2.6,4.8,5.2,4.4]
                if len(header)==3:widths=[3,7,7]
                for cell,width in zip(table.rows[0].cells,widths):cell.width=Cm(width)
                for cell,label in zip(table.rows[0].cells,header):
                    cell.text=label;cell.paragraphs[0].runs[0].bold=True
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFEFEF');cell._tc.get_or_add_tcPr().append(shade)
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for values in rows:
                    row=table.add_row()
                    for cell,width,value in zip(row.cells,widths,values):cell.width=Cm(width);inline(cell.paragraphs[0],value);covered.append(value)
                    nosplit=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nosplit)
                for row_index,row in enumerate(table.rows):
                    for c in row.cells:
                        for p in c.paragraphs:
                            p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.08
                            p.paragraph_format.keep_with_next=row_index<len(table.rows)-1
                            for r in p.runs:r.font.size=Pt(9.5)
                doc.add_paragraph()
            continue
        if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('### '):doc.add_heading(line[4:],2)
        elif line.startswith('## '):doc.add_heading(line[3:],1)
        elif re.match(r'^\d+\. ',line):inline(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
        elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        else:inline(doc.add_paragraph(),line)
        i+=1
    target=out/(title+'_V1.0.docx');doc.core_properties.title=title;doc.core_properties.subject='R2操作手册 同版Word阅读导出';doc.save(target)
    loaded=Document(target);parts=[p.text for p in loaded.paragraphs]+[c.text for t in loaded.tables for r in t.rows for c in r.cells]
    full='\n'.join(parts)
    # Table body content is plain; explicit hyperlinks may not appear in python-docx p.text.
    assert all(re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',v) in full for v in covered)
    manifest.append({'title':title,'source':str(src),'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'md':str(manual/'MD同版导出'/(title+'.md')),'docx':str(target),'tables':len(loaded.tables),'paragraphs':len(loaded.paragraphs)})
(work/'manual-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest,ensure_ascii=False,indent=2))
