from pathlib import Path
import re, zipfile, hashlib, json
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'E:\AI漫剧\作品集\龙凤胎长出蛇尾')
skill=Path(r'E:\AI漫剧\02_制作方法\02_专业技能\自研Skill\shot-language-design')
out=root/'第一集_1-1日间寝殿_文字分镜_V0.2_重设计待审.docx'
assert not out.exists(), 'Do not overwrite an existing review version'
protected=[root/'龙凤胎长出蛇尾前十集_208928.docx',root/'第一集项目剧本分析与导演本_导演讨论稿_V0.2.docx',root/'第一集_1-1日间寝殿_文字分镜_V0.1_待审.docx',root/'第一集_1-1日间寝殿_文字分镜_V0.1_待审.md']
before={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}
fields=['镜号／时间','镜头任务','起幅与构图','主体行动与场景动态','摄影机与光学','声音与对白','落幅、切点与承接']
dur=[1,5,2,4,3,4,6,2,10,3,5]
spoken=['','','咱家九代单传，','你一下添了两个，往后可有福享了！','奶奶，这太贵重了。','今晚是昭儿和灵儿的满月洗礼，','作为宗主夫人，你自然该拿着玉牌主理大小事务。','','浅浅，后山黑水玄蛇的封印有异，我去一趟，赶在两个孩子洗礼前回来。','嗯，路上小心。','']
# Six fields per shot; timings and identifiers are added from the same data.
body=[
['以金锁的亮点直接抓住注意，下一镜立即交代它属于谁。',
'一枚金锁与红绸占据画面，锁在原文的红绸位置，不佩在婴儿颈上。雕纹可读，不要求看清文字。',
'锁和红绸不自行漂动。通过视点变化看到真实金属高光移动；室内暖光有稳定来源。',
'金锁特写，镜头沿摇篮外缘短横移掠过锁面，红绸与后方篮沿产生视差；不转动金锁来假冒运镜。焦点守锁面。',
'安静室内底声；无对白、哭声。地点字幕留给下一镜，避免一秒内同时读锁与字。',
'金锁轮廓读到即切，不停留展示全套雕纹。以同一摇篮的篮沿位置接S02，近物切到所处关系。'],
['从孩子拉开到全家，尽早建立双婴及受珍视的空间关系。',
'起于双摇篮近中景，初始观察方向男婴左、女婴右，双脸均可辨；祖母和母亲先在画外。',
'婴儿睡眠呼吸带动局部衣褥，篮不摇。镜头拉开后看见祖母、母亲近距离对坐，凌云在母亲肩侧；祖母准备递牌但尚未完成交接。',
'摄影机真实升高并后拉，从斜俯双婴展开为室内全景；高度变化增加对座位与通路的可见范围，篮沿、人物、门框形成层次视差。走摇篮外侧净空，不越过婴儿正上方或穿椅；终点仍保交流轴观察侧。',
'地点字幕“天剑宗 静心苑”在前段出现，再于双婴仍可辨时对应标“凌昭”“凌灵”；识别不够时延长本镜预算，不把全部身份字幕叠在一瞬。无对白。',
'以祖母抬手向母亲递牌启动下一镜；动作在切点保持进行中。全景交代座位和通路后立刻切手部，避免巡视豪华空殿。'],
['用短时信息特写落实玉牌的归属变化。',
'两手与玉牌的特写：母亲承接掌心向上，祖母手从原来的交流方向进入；衣袖颜色和手的朝向延续S02。',
'祖母将玉牌按进母亲掌心，她接稳；接触令皮肤、袖口产生小幅挤压与回落。此后不再递第二次。',
'短稳读特写，机位略俯双手；摄影机此刻不动，留出看清触碰和交付的时间。清晰范围覆盖玉牌与接触处。',
'太御夫人画外对苏浅浅：“咱家九代单传，”语流在下镜继续，同一次发言。衣料细声低于台词。',
'玉牌已接稳，话音不停，切祖母的脸接余句；不用再次抬牌制造重复。'],
['让观众近距离接到祖母的喜悦和认可。',
'祖母近景，轻带母亲肩部作为关系锚点；视线对母亲，不朝镜头宣讲。',
'祖母笑着完成认可，眉眼与说话气息自然，手已从交牌处放松收回。没有转向女婴时的冷脸或反派暗示。',
'视线高度近景，稳读表情；前镜手部特写转人脸近景，尺度已有变化，不另加无作用缓推。',
'太御夫人接上镜：“你一下添了两个，往后可有福享了！”字幕“太御夫人”放人物无遮挡处。',
'祖母认可落到对方后切母亲接收，视线方向成对；不用祖母说完再完整重播一句。'],
['从被给予转到母亲接收重礼的情绪。',
'母亲近景，祖母肩旁同轴侧反看，面部和抬眼方向清楚，玉牌仍在母亲手里、可在下裁切外。',
'她从掌中玉牌抬眼向祖母，局促里有接收照顾的柔顺；不推回玉牌，不带隐瞒身世的惊惧。',
'短推近，由含肩近景收至更紧的脸部近景，随她抬眼进入回应；终点不截断下颌和眼神，背景透视随前移改变。不是数码放大。',
'苏浅浅对太御夫人：“奶奶，这太贵重了。”字幕“宗主夫人 苏浅浅”。',
'回应结束，凌云的手进入肩侧动作接下一镜衣料接触；切点接行动，不增加一轮无信息对望。'],
['把丈夫的体贴拍成具体触碰，给长句建立可看的行动。',
'肩部衣料与凌云手的特写，保留墨蓝袖口表明归属；苏浅浅持牌手不参与拢衣。',
'凌云从肩侧提拢滑开的披衣，布料随受力收拢，有拉力、褶皱聚集与落回，不像手划过衣服却毫无反馈。',
'摄影机沿肩线侧移跟随拢衣的手，始终让接触点落在清楚区域；路径走人物前侧的外缘，不穿肩颈。侧移造成衣褶与后景的层次变化。',
'凌云画外开始：“今晚是昭儿和灵儿的满月洗礼，”本句在S07继续，不把镜头切换当成重新说话。',
'手完成最后收拢尚未完全收回时切S07同一动作末段；下一镜只接结果，不重新披一次衣。'],
['由触碰展开到夫妻关系，让身份认可落在对方身上。',
'夫妻近中景，母亲坐、丈夫在肩侧俯身；两人的脸和彼此朝向可辨，玉牌已在她手中。',
'丈夫接着说话、收回拢衣手，妻子以抬眼和放松肩部接收。披衣保持收好状态，衣袖随收手落下。',
'沿夫妻前侧做短弧绕行，绕行中心为两人关系区域，由斜侧双人近中景转到更靠丈夫正面的近景，妻子侧脸留边缘。摄影机在地面空间走弧线，人物不随镜头转身；不滚转画面、不越交流轴，弧度服从脸部遮挡和净空。',
'凌云续：“作为宗主夫人，你自然该拿着玉牌主理大小事务。”字幕“天剑宗宗主 凌云”。句完山外钟声陡入，凌云注意转向腰间。',
'钟声令注意从妻子转开即切巡山令，接同一注意动作；允许运动中切，不要求每次运镜先停稳再等切。'],
['用紧凑的物件警讯打断照顾节奏。',
'腰间巡山令特写，墨蓝衣料确认属于凌云；令不被拿到手里。',
'同一次钟声下令牌亮蓝光，只照到腰际近处衣料；衣料延续他转注意时的小幅牵动，不全殿爆闪。',
'由稍松的腰际局部做短促推进，压到令牌特写后留可读停顿；速度比前段舒缓展开明显更利落。路径在腰前净空，不能穿进令牌或身体。',
'S07钟声余响延续，不再另敲一记配合发光；无令牌语音、蛇嘶或新增对白。',
'蓝光读到后切凌云已经注意到警讯的近景，不重复先无反应再重新发现。'],
['从个人接警转为夫妻交代，兼容原因、离开与回归承诺。',
'凌云近景，目光从腰际方向回到妻子；太御夫人仍在原座画外。',
'他皱眉说明事务，向妻子交代；妻子听完，不因“蛇”字惊恐。说到回来时两人视线接住，暂不转身离场。',
'起于凌云脸部近景，随他转回妻子开始后拉，逐渐纳入坐着的妻子，落成有距离关系的双人中景；后半句保持清楚可读。后拉让职责和被留下的家庭同框，不持续匀速拉到人脸不可辨。',
'凌云对苏浅浅：“浅浅，后山黑水玄蛇的封印有异，我去一趟，赶在两个孩子洗礼前回来。”钟声渐退，完整话语只发声一次。',
'承诺落在妻子身上后切她近景，拉开不是情感疏远暗示；不切殿外蛇妖或凭空出现的战斗。'],
['给母亲的关心一个清楚、短促的反应落点。',
'母亲近景，同观察侧看向站立丈夫，视线略上扬；玉牌未转交他人。',
'母亲点头说路上小心，真实呼吸与眼神回应即可，不追加欲言又止或藏秘密的惊恐。',
'近景短稳读；对比上镜后拉与下镜离场运动，保住这一句真切的关心。',
'苏浅浅对凌云：“嗯，路上小心。”丈夫在这句结束后才转身，不抢走她的回应。',
'丈夫起身转向的衣料声可从尾部进入，切S11接同一转身离去。'],
['用人物离开与画面展开收束白天，保留家庭的完整末态。',
'双人中景接凌云转身，母亲坐在原位；门位于既定通路尽头，两摇篮仍在内室，不挡离开线路。',
'凌云沿通路走向门口，衣摆随步伐摆动，脚步反射符合室内距离。母亲不追出；祖母与双婴留在室内正常状态。',
'摄影机沿落座区外缘后移拉开，由双人中景展开为包含母亲、摇篮及门口的关系全景；凌云向门走，摄影机与他运动方向不同，拉开双方距离。通路进入景深后让他出画，镜头停在留下的家人，不跟去殿外。',
'无新增对白；脚步和衣料随远离减弱，室内底声连续。此时不加孩子哭声预告夜间。',
'丈夫离开、家庭末态读清即收；1-2是时间省略后的夜景，不把同几步脚声硬接到深夜。']
]
assert len(body)==len(dur)==len(spoken)==11
shots=[];start=0
for i,(row,d) in enumerate(zip(body,dur)):
    assert len(row)==6
    shots.append([f'E01-S{i+1:02}｜00:{start:02}—00:{start+d:02}｜{d}秒']+row);start+=d
intro=[
'本稿覆盖第一集1-1日间寝殿。V0.1已被用户否定，保留作历史对照，不作为采用稿。本稿V0.2重新提交审阅，承接导演本V0.2的D1意图，暂不进入1-2夜场。',
'设计主线是先让观众看见被珍视的双婴，再读清赠牌与拢衣的触碰，由钟声和蓝光压出一次节奏变化，最后回到夫妻交代与离开。对白可以跨镜持续，切镜跟着信息和行动转移。',
'景别骨架：金锁特写 → 双婴近中景升拉至全景 → 赠牌手部特写 → 祖母近景 → 母亲推近 → 拢衣特写横移 → 夫妻短弧绕行近景 → 令牌特写短促推入 → 丈夫近景后拉成双人中景 → 母亲近景 → 离场后拉成关系全景。全景交代关系，近景读情绪，特写给短时信息，不按种类凑数。',
'空间为本轮布置提案：两摇篮并列，母亲与祖母在同端近距离对坐，递牌不跨摇篮；凌云在母亲肩侧，出口由落座区外缘通路可达。主要观察在祖母—母亲交流轴靠摇篮侧。实际净空与遮挡仍须随场景资产确认。婴儿左右以S02初始观察为准，世界位置不随换机位调换。',
'共性：AI仿真人竖屏，具体比例、分辨率与素材未锁；日间暖而可读，材质与光源连续。摄影机位移产生视差，衣料变化来自接触和步伐，环境不随机晃动。人物身份字幕分次出现；字幕可读时间须随实际尺寸复核。',
'11镜预算45秒，各镜1、5、2、4、3、4、6、2、10、3、5秒，是设计预估，不是试读或动态预演结果。长句须先试读再调，不能为锁时码吞字。全集90—150秒要求继续有效，夜场预算另核；本场不承担全集全部高潮。'
]
checks=[
'原文覆盖：金锁与双婴S01—02；赠牌S03；祖母认可S03—04；母亲回应S05；拢衣和丈夫认可S06—07；钟声/蓝光S07尾—08；离开原因与回归承诺S09；妻子关心S10；离场S11。五句原对白完整、顺序不变，祖母第一句及丈夫认可句分两镜连续发声，无重复。',
'静止阅读点是赠牌接触S03、祖母喜悦S04和母亲回应S10；其余运动分别负责揭空间、近情绪、跟触碰、绕关系、压警讯、扩交流与离场。不是以运动数量证明质量，后续须在预演中看路径和速度是否成立。',
'持物与状态：S03之后玉牌归母亲，特写裁切外也不得重置；S06完成拢衣后不回到散开状态；巡山令仍佩腰间；两婴白天无异变。没有新增身世揭示、敌人、男婴蛇尾、母亲追出或黑化表情。',
'编号变化：旧S01→新S01—02；旧S02→新S03—04；旧S03→新S05；旧S04→新S06—07；旧S05→新S08；旧S06→新S09；旧S07→新S10—11。新编号只属于V0.2，引用旧稿时必须带版本。',
'Markdown是七列表，Word由同一数据转为七项竖排阅读版，内容相同。当前只做文字核查和文档排版检查，没有生成媒体，未做配音测时，未声称用户采用。下一起点仍为1-2月光落到摇篮。'
]
with zipfile.ZipFile(protected[0]) as z:
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    a=[''.join(x.xpath('.//w:t/text()',namespaces=ns)) for x in etree.fromstring(z.read('word/document.xml')).xpath('//w:body//w:p',namespaces=ns)]
original=[]
for text in a[21:36]:
    m=re.match(r'^(太御夫人|苏浅浅|凌云)(?:[（(][^）)]*[）)])?\s*[：:]\s*(.*)$',text)
    if m: original.append(m.group(2))
assert len(original)==5
assert ''.join(original)==''.join(spoken)
assert all((not s) or s in body[i][4] for i,s in enumerate(spoken))
assert start==45
md='# 第一集日间寝殿文字分镜\n\nV0.2｜2026-10-27｜重设计待审｜shot-language-design V0.11\n\n'+'\n\n'.join(intro)+'\n\n'
md+='| '+' | '.join(fields)+' |\n|'+'|'.join(['---']*7)+'|\n'
for row in shots: md+='| '+' | '.join(x.replace('|','／') for x in row)+' |\n'
md+='\n## 回查与版本说明\n\n'+'\n\n'.join(checks)+'\n'
out.with_suffix('.md').write_text(md,encoding='utf-8')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor(0,0,0)
    for border in list(st.element.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Title'].font.size=Pt(21)
for name in ['Heading 1','Heading 2']:doc.styles[name].font.size=Pt(14 if name=='Heading 1' else 12)
sec.header.paragraphs[0].text='龙凤胎长出蛇尾 第一集 日间寝殿 V0.2重设计待审';sec.header.paragraphs[0].runs[0].font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');sec.footer.paragraphs[0]._p.append(f)
doc.add_paragraph('第一集日间寝殿文字分镜',style='Title')
doc.add_paragraph('十一镜  设计预算45秒  重设计待审')
doc.add_paragraph('2026年10月27日  镜头语言设计Skill V0.11')
for t in intro:doc.add_paragraph(t)
for i,row in enumerate(shots):
    if i%2==0:
        p=doc.add_heading('文字分镜 '+str(i+1)+((' 至 '+str(i+2)) if i+1<len(shots) else ''),1);p.paragraph_format.page_break_before=True
    doc.add_heading(row[0].replace('｜','  '),2)
    for label,value in zip(fields[1:],row[1:]):
        p=doc.add_paragraph();p.add_run(label+'：').bold=True;p.add_run(value)
p=doc.add_heading('回查与版本说明',1);p.paragraph_format.page_break_before=True
for t in checks:doc.add_paragraph(t)
doc.core_properties.title='龙凤胎长出蛇尾 第一集日间寝殿文字分镜'
doc.save(out)
text='\n'.join(p.text for p in Document(out).paragraphs)
assert all(value in text for row in shots for value in row[1:])
assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in before.items())
Path(__file__).with_name('重设计核查.json').write_text(json.dumps({'shots':11,'estimated_seconds':45,'original_dialogue_complete_and_ordered':True,'protected_files_unchanged':before,'new_docx':str(out)},ensure_ascii=False,indent=2),encoding='utf-8')
print(str(out));print('11 shots; 45 seconds estimated; source dialogue exact; old files and sources unchanged')
