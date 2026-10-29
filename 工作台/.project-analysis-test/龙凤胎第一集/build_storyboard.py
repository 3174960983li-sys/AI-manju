from pathlib import Path
import json,re,zipfile,hashlib
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'E:\AI漫剧\作品集\龙凤胎长出蛇尾')
src=root/'龙凤胎长出蛇尾前十集_208928.docx'
director=root/'第一集项目剧本分析与导演本_导演讨论稿_V0.2.docx'
out=root/'第一集_1-1日间寝殿_文字分镜_V0.1_待审.docx'
mdpath=out.with_suffix('.md')
assert out.name=='第一集_1-1日间寝殿_文字分镜_V0.1_待审.docx'
before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [src,director]}

fields=['镜号／时间','镜头任务','起幅与构图','主体行动与场景动态','摄影机与光学','声音与对白','落幅、切点与承接']
shots=[
['E01-S01｜00:00—00:05｜5秒',
'先认清两个满月婴儿及被珍视的处境，直接进入故事，不先巡视空殿。',
'双摇篮并排，当前观察男婴在左、女婴在右；红绸与两枚长命锁处于原文所写位置，两个婴儿的脸均可辨。',
'婴儿熟睡，只有自然呼吸带动局部衣褥，摇篮不自行晃动。长命锁压在红绸处，不改成戴在孩子身上。',
'高位斜俯双摇篮中景，固定。以足够清晰范围保住两婴及锁的关系，避开夸张广角和过浅景深。',
'室内轻微底声，无新增对白或婴儿哭声。字幕先“天剑宗 静心苑”，随后对应两婴显示“凌昭”“凌灵”，错开出现避免堆叠；可读时间待实际画幅检验。',
'两婴身份和金锁关系读清后切赠牌。下镜从同一家庭场面的人物交互承接，不重排摇篮身份。'],
['E01-S02｜00:05—00:11｜6秒',
'把“受宠”落实成祖母的交付与认可，而不只交代豪华陈设。',
'从交流轴同一侧斜看近距离对坐的两人：太御夫人在画面左，苏浅浅在右；交牌双手位于两人之间可辨区域，不被摇篮遮挡。',
'太御夫人将玉牌按入苏浅浅承接的掌心，再完成原句；苏浅浅接稳，不急着推回。衣袖只随手部接触移动，两人不为配合对白反复递牌。',
'固定斜侧双人中景，机位略高于坐姿视线、轻俯以包含手与面部。清晰关系覆盖交接和祖母的表达，不以手部极近景割断交流。',
'太御夫人，画内对苏浅浅：“咱家九代单传，你一下添了两个，往后可有福享了！”音量亲近。字幕“太御夫人”随人物清楚出现；室内底声连续。',
'玉牌已归苏浅浅，祖母完成话语后切她接收重礼的回应。下一镜不得再次交牌，衣袖和持牌手状态承接。'],
['E01-S03｜00:11—00:15｜4秒',
'读到苏浅浅面对重礼的局促，给凌云的认可建立回应对象。',
'同一轴侧、祖母肩旁看苏浅浅三分之二侧脸；她的视线仍向祖母。玉牌由她持有，手可在下裁切外，但不发生新交接。',
'她先看已经接住的玉牌，再抬眼向祖母说话，回应中有受宠的迟疑。不是盯着孩子露出心虚，也不退还玉牌。',
'坐姿视线高度的近景，固定；焦点优先面部。祖母肩部作为轻前景，不遮口型；不要求近景同时看清膝上玉牌。',
'苏浅浅，画内对太御夫人：“奶奶，这太贵重了。”字幕“宗主夫人 苏浅浅”在本镜出现，位置避开面部；无新增OS。',
'“贵重了”及她望向祖母的回应成立后，接凌云在肩侧的照顾。手仍持牌，下一镜扩大范围后重新看见其稳定归属。'],
['E01-S04｜00:15—00:26｜11秒',
'在同一连续照顾过程中，建立夫妻温柔和对她身份的认可。',
'坐站双人中景：苏浅浅居画面中下，凌云在她肩侧上方；构图保留她持玉牌的手与他拢衣的手，二者不被衣料遮没。祖母仍在原座，可处画外。',
'凌云俯身替她拢好肩侧披衣，朝她说话；她听着，接收认可，继续持牌。拢衣与对白重叠，不额外增加拥吻或全员反应。衣褶随接触收拢并留住结果。',
'机位留在祖母—母亲交流轴靠摇篮的一侧；随凌云进入照顾动作，沿落座区外缘短距离横移，避开肩部遮挡，停在两张脸、拢衣和持牌关系都可读的位置。全过程不跨轴、不穿座椅；中等透视、足够清晰层次。',
'凌云，画内对苏浅浅：“今晚是昭儿和灵儿的满月洗礼，作为宗主夫人，你自然该拿着玉牌主理大小事务。”字幕“天剑宗宗主 凌云”。话后、拢衣已完成时，第一下山外钟声进入；不提前盖住对白。',
'钟声令凌云注意转变时切腰间巡山令；同一次钟声跨镜延响，衣已拢好，玉牌未离苏浅浅。机位短横移在切前已稳定。'],
['E01-S05｜00:26—00:29｜3秒',
'让宗门警讯有可见载体，说明随后转向事务的触发。',
'凌云腰间巡山令与墨蓝衣料的局部，令牌仍佩在腰间，没有被拿到手中。',
'巡山令在同一声钟的进入阶段应声亮蓝光，近处衣料获得有限反光；不扩大成全殿闪烁。拢衣的手已收回，不从衣领瞬移到握牌。',
'腰部高度的令牌近景，固定，从人物原观察侧取局部。焦点落在巡山令，保留腰际衣料作归属锚点；不将此蓝光做成妖物视角。',
'承接S04第一次钟声的尾响，不重新敲一次来配画面。无对白、无令牌语音，也不添加蛇嘶声。',
'蓝光警讯读清后切凌云；下一镜从他注意腰间的状态接向抬眼对妻子交代，不重复亮牌事件。'],
['E01-S06｜00:29—00:41｜12秒',
'交代离开的原因与会回来的承诺，家庭关系保持信任。',
'同轴侧看凌云与坐着的苏浅浅；他先注意令牌方向，再看向她，腰际蓝光可留作轻微归属提示。她持牌于原位置。',
'凌云从警讯方向抬眼，带职责被触发的皱眉朝妻子说明情况；说完准备转向出口，不先离开再向无人处交代。苏浅浅在同一交流中听完，不因蛇妖名称突然惊恐。',
'固定坐站双人中景偏凌云，机位在坐站之间的合适高度，使双方视线和身份可辨。清晰层次同时照顾说者、听者，不在长句中无因推拉或转焦。',
'凌云，画内对苏浅浅：“浅浅，后山黑水玄蛇的封印有异，我去一趟，赶在两个孩子洗礼前回来。”钟声尾响自然消退，听感不抢句中信息。',
'交代完且妻子已接收到信息后切稍宽关系。她的答复只在S07出现；不新增黑水玄蛇、禁地或战斗画面。'],
['E01-S07｜00:41—00:46｜5秒',
'接住妻子的关心，让丈夫离开，交出可接月夜的家庭末态。',
'稍宽的纵深关系：苏浅浅近处坐着，凌云在肩侧，门与不穿过摇篮的离开通路在两人后方偏母亲一侧。原位祖母可留边缘或画外。',
'苏浅浅点头回应，玉牌仍归她；凌云得到回应后转身沿通路离开画面，衣摆随步伐自然摆动。母亲不追出，孩子没有异变。',
'仍在同一交流轴侧的较宽中景，固定。用人物离开产生距离变化，不再跟去殿外；景深使母亲和离去动作可辨，不强求门外细节。',
'苏浅浅，画内对凌云：“嗯，路上小心。”随后衣料与脚步随离开延续、减弱，室内底声保留。',
'凌云已离开，母亲与祖母及双婴留在室内；玉牌、披衣和摇篮位置不重置。下段1-2是经过时间省略的夜景，允许切换声场，不将这几步脚声无依据接到深夜。']
]

intro=[
'覆盖：第一集1-1日间寝殿全场。依据原稿与导演本V0.2的D1；用户已要求继续该方案的镜头设计，沿用其核心意图，不要求再次采用同一上游。当前文字分镜V0.1待审，夜间1-2尚未展开；下一起点为月光落到摇篮。',
'目标：让观众通过双婴、赠牌、拢衣读到珍视，再经钟声和巡山令转向事务，结束于妻子关心、丈夫离开。五句本场原对白完整保留，无新增OS、蛇尾或身份揭底。',
'本轮布置提案：两摇篮并列；母亲与太御夫人在摇篮同端的近距离落座区对坐，交牌不跨摇篮。主观察关系中祖母左、母亲右，凌云在母亲肩侧，出口在两人后方偏母亲侧；主要交流机位均留在祖母—母亲轴靠摇篮一侧。具体尺度与通路须随实际场景资产复核，不冒称原文已锁。S01的婴儿左右对应依该镜观察保留，不能在后续反向观察时调换世界位置。',
'共同基准：AI仿真人、竖屏，具体比例/分辨率未锁；日间暖而清楚的受光，光源与衣料材质连续。空间、衣褶、玉牌接触和脚步提供真实变化，不添加无因风雨或随机镜头抖动。人物/服装/场景参考尚未提供，不编素材号。',
'时间是逐镜设计预算：5＋6＋4＋11＋3＋12＋5＝46秒。它承接导演日场约45秒的预算，差1秒属于待试读校准的估计调整；全集仍以90—150秒为约束。未测语速、未做动态预演，不保证逐秒执行，也不强迫生成任务按七镜一一拆分。'
]

coverage=[
'原文→镜头：双摇篮、红绸、两枚长命锁及双婴身份在S01；太御夫人赠牌、原话与认可在S02；苏浅浅接礼的回应在S03；拢衣和宗主夫人身份认可在S04；钟声/巡山令警讯在S04尾—S05；离开原因与回归承诺在S06；妻子关心及离开在S07。',
'台词与声音：本场五句原对白各在一镜完整发声，不跨镜重复；夜场第六句尚未进入。S04尾进入的山外钟声延续S05，是同一个故事事件，不是重新起钟。角色能听到钟声，不对配乐作反应。',
'状态：S02交出的玉牌自此归苏浅浅；S03近景允许手在裁切外，S04再次包含稳定持牌，S07不消失。披衣在S04整理后保持；巡山令仍佩腰间。双婴在本段正常，女婴夜间异变另段处理。',
'运镜回查：S04是唯一主动横移，必须用落座区外缘可通行的路径避开遮挡并停稳，不用运动术语掩盖穿人/穿椅。S01与其他固定镜头靠人物、物件和声音推进，固定不等于环境冻结。若实际布景不支持这条短路径，先调整未锁机位/布置，不能为拍摄无因新增通路。',
'阅读方式：Markdown保留七列表；Word由同一组字段生成，每镜采用七项竖排，避免在竖向纸面缩成难读宽表。两份内容同版，修改以同一镜号同步，不另建第二套动作。',
'状态与下一步：本轮止于日间段完整文字设计，待用户反馈后再接1-2夜间段。未生成故事板、TIG图、关键帧、配音或视频提示词；没有实际媒体可读性或精确时长的通过声明。'
]

# Check original dialogue once per audio field, not once across every document reference.
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(src) as z:
    a=[''.join(x.xpath('.//w:t/text()',namespaces=ns)) for x in etree.fromstring(z.read('word/document.xml')).xpath('//w:body//w:p',namespaces=ns)]
original=[]
for text in a[21:36]:
    m=re.match(r'^(太御夫人|苏浅浅|凌云)(?:[（(][^）)]*[）)])?\s*[：:]\s*(.*)$',text)
    if m:original.append(m.group(2))
assert len(original)==5
audio='\n'.join(row[5] for row in shots)
for line in original:assert audio.count(line)==1,('Missing or repeated dialogue',line)
assert sum([5,6,4,11,3,12,5])==46
assert all(len(row)==7 for row in shots)

markdown='# 第一集日间寝殿文字分镜\n\nV0.1｜2026-10-27｜待审｜shot-language-design V0.10\n\n'
markdown+='\n\n'.join(intro)+'\n\n'
markdown+='| '+' | '.join(fields)+' |\n|'+ '|'.join(['---']*7)+'|\n'
for row in shots:markdown+='| '+' | '.join(s.replace('|','／').replace('\n','<br>') for s in row)+' |\n'
markdown+='\n## 覆盖与连续性回查\n\n'+'\n\n'.join(coverage)+'\n'
mdpath.write_text(markdown,encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor(0,0,0)
    for border in list(st.element.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.12;doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Title'].font.size=Pt(21);doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Subtitle'].font.italic=False
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.size=Pt(14 if name=='Heading 1' else 12)
    doc.styles[name].paragraph_format.space_before=Pt(8);doc.styles[name].paragraph_format.space_after=Pt(5)
sec.header.paragraphs[0].text='龙凤胎长出蛇尾  第一集  日间寝殿文字分镜  V0.1待审'
sec.header.paragraphs[0].runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('第 ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f);foot.add_run(' 页')
doc.add_paragraph('第一集日间寝殿文字分镜',style='Title')
doc.add_paragraph('1-1全场  七镜  设计预算46秒  待审',style='Subtitle')
doc.add_paragraph('2026年10月27日  镜头语言设计Skill V0.10')
doc.add_heading('段落设计说明',1)
for t in intro:doc.add_paragraph(t)
doc.add_heading('镜头组织的依据',1)
doc.add_paragraph('从婴儿与金锁起步，赠牌使认可具体，妻子的回应引出丈夫照顾；同镜横移服务拢衣与交流，令牌插镜切换注意，再回到夫妻交代和离开。每镜承担主要任务，不按一句台词必切一次或每镜必须运镜的模板设计。')

for i,row in enumerate(shots):
    if i%2==0:
        label='文字分镜 '+str(i+1)+((' 至 '+str(i+2)) if i+1<len(shots) else '')
        par=doc.add_heading(label,1);par.paragraph_format.page_break_before=True
    doc.add_heading(row[0].replace('｜','  '),2)
    for name,text in zip(fields[1:],row[1:]):
        par=doc.add_paragraph();par.add_run(name+'：').bold=True;par.add_run(text)

doc.add_heading('段尾状态与回查',1)
for text in coverage[:4]:doc.add_paragraph(text)
doc.add_paragraph('本段待审；下一起点为1-2月光落到摇篮。尚未进入图像或提示词分支。原文与导演本未修改。')
doc.core_properties.title='龙凤胎长出蛇尾 第一集日间寝殿文字分镜'
doc.core_properties.subject='1-1全场 E01-S01至E01-S07 七项竖排阅读版'
doc.save(out)
assert all(hashlib.sha256((root/name).read_bytes()).hexdigest()==h for name,h in before.items())
check=Document(out)
alltext='\n'.join(p.text for p in check.paragraphs)
for row in shots:
    for field in row[1:]:assert field in alltext
print(out)
print(mdpath)
print('7 shots, 46 seconds design budget, all 5 source dialogue lines once in audio fields, sources unchanged')
